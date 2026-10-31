"""TEST: openpyxl — ValueError from an invalid cell reference."""

import pytest
from openpyxl import Workbook
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.exceptions import IllegalCharacterError


def test_openpyxl_invalid_cell_reference():
    """Setting an invalid cell reference raises ValueError."""
    wb = Workbook()
    ws = wb.active
    with pytest.raises(ValueError):
        ws["ZZZZ99999"] = "Hello"


def test_openpyxl_valid_cell_reference():
    """The correct version: use a valid cell reference like 'A1'."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Hello"
    assert ws["A1"].value == "Hello"
